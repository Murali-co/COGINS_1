import io
import fitz  # PyMuPDF
import docx  # python-docx

class ResumeParser:
    @staticmethod
    def parse_pdf(file_bytes: bytes) -> str:
        text = ""
        try:
            # Try PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page in doc:
                text += page.get_text()
            doc.close()
        except Exception as e:
            # Fallback to pdfminer if available
            print(f"PyMuPDF failed, falling back to pdfminer: {e}")
            text = ""

        if not text.strip():
            try:
                from pdfminer.high_level import extract_text as extract_pdfminer_text

                fp = io.BytesIO(file_bytes)
                text = extract_pdfminer_text(fp)
            except ModuleNotFoundError:
                print("pdfminer is not installed; PDF fallback extraction is unavailable.")
            except Exception as e:
                print(f"pdfminer fallback failed: {e}")

        return text

    @staticmethod
    def parse_docx(file_bytes: bytes) -> str:
        text = []
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text.append(para.text)
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text.append(" | ".join(row_text))
        except Exception as e:
            print(f"python-docx parsing failed: {e}")
            
        return "\n".join(text)

    @classmethod
    def parse(cls, file_bytes: bytes, filename: str) -> str:
        lower_filename = filename.lower()
        if lower_filename.endswith(".pdf"):
            return cls.parse_pdf(file_bytes)
        elif lower_filename.endswith(".docx"):
            return cls.parse_docx(file_bytes)
        else:
            raise ValueError("Unsupported file format. Only PDF and DOCX are allowed.")
